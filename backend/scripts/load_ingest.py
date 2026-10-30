"""100 concurrent POST /ingest gate (suggested-order item 31 / Docs/specs/16-load.md).

Setup once: tiny .docx → presign → S3 PUT. Then fire LOAD_INGEST_CONCURRENCY
parallel ingest calls with unique jobIds against the shared object key.

Requires a live FastAPI with TB_ASSERTION_KEYS and TB_S3_* matching this process.
Leave AI_API_KEY empty on the server so the run does not call an LLM.
"""

from __future__ import annotations

import asyncio
import os
import sys
import time
import uuid
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import httpx
from docx import Document

# Make auth_py + backend importable when run as `uv run python scripts/load_ingest.py`
_ROOT = Path(__file__).resolve().parents[1]
_AUTH = _ROOT.parent / "auth" / "py" / "src"
sys.path.insert(0, str(_ROOT / "src"))
sys.path.insert(0, str(_AUTH))

from auth_py.assertion import (  # noqa: E402
    ASSERTION_AUDIENCE,
    ASSERTION_TTL_SECONDS,
    AssertionClaims,
    mint_assertion,
    parse_assertion_keys,
)

BASE = os.environ.get("LOAD_INGEST_BASE_URL", "http://127.0.0.1:8080").rstrip("/")
CONCURRENCY = int(os.environ.get("LOAD_INGEST_CONCURRENCY", "100"))
TIMEOUT_S = float(os.environ.get("LOAD_INGEST_TIMEOUT_S", "120"))
MAX_WALL_S = float(os.environ.get("LOAD_INGEST_MAX_WALL_S", "180"))
UID = os.environ.get("LOAD_INGEST_UID", "load-user")
ROLE = os.environ.get("LOAD_INGEST_ROLE", "author")
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _keys() -> tuple[str, bytes]:
    raw = os.environ.get("TB_ASSERTION_KEYS", "").strip()
    if not raw:
        # Load backend/.env if present (dev convenience; never commit secrets).
        env_path = _ROOT / ".env"
        if env_path.is_file():
            for line in env_path.read_text(encoding="utf-8").splitlines():
                if line.startswith("TB_ASSERTION_KEYS="):
                    raw = line.split("=", 1)[1].strip().strip('"').strip("'")
                    break
    parsed = parse_assertion_keys(raw)
    if not parsed:
        raise SystemExit("TB_ASSERTION_KEYS is required (kid:secret matching the running service).")
    kid, secret = next(iter(parsed.items()))
    return kid, secret


def _token(kid: str, secret: bytes) -> str:
    now = int(datetime.now(UTC).timestamp())
    return mint_assertion(
        AssertionClaims(
            aud=ASSERTION_AUDIENCE,
            exp=now + ASSERTION_TTL_SECONDS,
            gen=1,
            iat=now,
            kid=kid,
            role=ROLE,
            sub="load-identity",
            uid=UID,
        ),
        secret,
    )


def _docx_bytes() -> bytes:
    document = Document()
    document.core_properties.title = "Load Test Runbook"
    document.add_paragraph("Step 1. Confirm the host is reachable.")
    document.add_paragraph("Step 2. Restart the service and re-check.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _setup(client: httpx.AsyncClient, kid: str, secret: bytes) -> tuple[str, int]:
    data = _docx_bytes()
    headers = {"Authorization": f"Bearer {_token(kid, secret)}"}
    presign = await client.post(
        f"{BASE}/v1/uploads/presign",
        headers=headers,
        json={
            "fileName": "load-test.docx",
            "contentType": DOCX_TYPE,
            "sizeBytes": len(data),
        },
    )
    if presign.status_code != 200:
        raise SystemExit(f"presign failed: {presign.status_code} {presign.text}")
    body = presign.json()
    put = await client.put(
        body["url"],
        content=data,
        headers=body.get("headers") or {"Content-Type": DOCX_TYPE},
    )
    if put.status_code not in {200, 201, 204}:
        raise SystemExit(f"S3 PUT failed: {put.status_code} {put.text}")
    return body["key"], len(data)


async def _one(
    client: httpx.AsyncClient,
    kid: str,
    secret: bytes,
    s3_key: str,
    size: int,
    index: int,
) -> tuple[bool, float, str]:
    t0 = time.perf_counter()
    job_id = f"load-{index}-{uuid.uuid4().hex[:10]}"
    try:
        response = await client.post(
            f"{BASE}/ingest",
            headers={"Authorization": f"Bearer {_token(kid, secret)}"},
            json={
                "jobId": job_id,
                "s3Key": s3_key,
                "fileName": "load-test.docx",
                "sourceKind": "docx",
                "sizeBytes": size,
            },
            timeout=TIMEOUT_S,
        )
        ms = (time.perf_counter() - t0) * 1000
        if response.status_code != 200:
            return False, ms, f"{job_id} status={response.status_code} {response.text[:200]}"
        payload = response.json()
        if not payload.get("draft"):
            return False, ms, f"{job_id} missing draft"
        return True, ms, job_id
    except Exception as exc:
        ms = (time.perf_counter() - t0) * 1000
        return False, ms, f"{job_id} {type(exc).__name__}: {exc}"


async def main() -> None:
    kid, secret = _keys()
    print(
        f"Load ingest: {BASE} concurrency={CONCURRENCY} "
        f"timeout={TIMEOUT_S}s maxWall={MAX_WALL_S}s uid={UID}"
    )

    limits = httpx.Limits(max_connections=CONCURRENCY, max_keepalive_connections=CONCURRENCY)
    async with httpx.AsyncClient(limits=limits) as client:
        health = await client.get(f"{BASE}/healthz", timeout=10)
        if health.status_code != 200:
            raise SystemExit(f"healthz failed: {health.status_code}")

        s3_key, size = await _setup(client, kid, secret)
        print(f"fixture key={s3_key} bytes={size}")

        wall0 = time.perf_counter()
        results = await asyncio.gather(
            *[_one(client, kid, secret, s3_key, size, i) for i in range(CONCURRENCY)]
        )
        wall_s = time.perf_counter() - wall0

    oks = [ms for ok, ms, _ in results if ok]
    fails = [msg for ok, _, msg in results if not ok]
    lat = sorted(oks)
    p99 = lat[min(len(lat) - 1, max(0, int(len(lat) * 0.99) - 1))] if lat else 0

    summary = {
        "concurrency": CONCURRENCY,
        "ok": len(oks),
        "failed": len(fails),
        "wallS": round(wall_s, 2),
        "p99Ms": round(p99),
    }
    print(summary)
    if fails[:5]:
        print("sample failures:")
        for row in fails[:5]:
            print(" ", row)

    bad = []
    if len(oks) != CONCURRENCY:
        bad.append(f"ok {len(oks)}/{CONCURRENCY}")
    if wall_s > MAX_WALL_S:
        bad.append(f"wall {wall_s:.1f}s > {MAX_WALL_S}s")
    if bad:
        print("FAIL:", "; ".join(bad))
        raise SystemExit(1)
    print("PASS")


if __name__ == "__main__":
    asyncio.run(main())
