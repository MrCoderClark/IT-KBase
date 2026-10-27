"""Word documents become the five article sections without an LLM."""

from __future__ import annotations

import base64
from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from fastapi.testclient import TestClient

from backend.parse.docx import DocxError, parse_docx


def _docx() -> Document:
    document = Document()
    document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    return document


def _bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_headings_map_into_article_sections() -> None:
    document = _docx()
    document.core_properties.title = "Linux SSH Key Rotation"
    document.core_properties.subject = "Rotate keys on production hosts."
    document.add_heading("Overview", level=1)
    document.add_paragraph("Keys are rotated every 90 days.")
    document.add_heading("Prerequisites", level=1)
    document.add_paragraph("Root access", style="List Bullet")
    document.add_paragraph("ansible-core", style="List Bullet")
    document.add_heading("Procedure", level=1)
    document.add_paragraph("1. Open the host")
    document.add_paragraph("ssh admin@host", style="Code Block")
    document.add_paragraph("2. Generate a key")
    document.add_paragraph("Use ed25519.")
    document.add_heading("Verification", level=1)
    document.add_paragraph("ssh -T succeeds", style="List Bullet")
    document.add_heading("Troubleshooting", level=1)
    document.add_heading("Permission denied", level=2)
    document.add_paragraph("Check authorized_keys.")
    document.add_heading("Rollback notes", level=1)
    document.add_paragraph("Keep the previous key for 7 days.")

    draft = parse_docx(_bytes(document), "ignored.docx")

    assert draft.title == "Linux SSH Key Rotation"
    assert draft.summary == "Rotate keys on production hosts."
    assert "Keys are rotated every 90 days." in draft.overview
    assert "Rollback notes" in draft.overview
    assert "Keep the previous key for 7 days." in draft.overview
    assert draft.prerequisites == ["Root access", "ansible-core"]
    assert draft.steps[0].title == "Open the host"
    assert draft.steps[0].code == "ssh admin@host"
    assert draft.steps[1].title == "Generate a key"
    assert "ed25519" in draft.steps[1].body
    assert draft.verification == ["ssh -T succeeds"]
    assert draft.troubleshooting[0].title == "Permission denied"
    assert "authorized_keys" in draft.troubleshooting[0].body


_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def test_numbered_procedure_keeps_pictures_on_the_step() -> None:
    document = Document()
    document.core_properties.title = "Access Onet client files"
    document.add_paragraph("Log into Microsoft OneDrive", style="List Number")
    document.add_paragraph("Use the company account.")
    document.add_picture(BytesIO(_PNG))
    document.add_paragraph("Open the recycle bin on the left", style="List Number")

    draft = parse_docx(_bytes(document), "onet.docx")

    assert draft.steps[0].title == "Log into Microsoft OneDrive"
    assert "company account" in draft.steps[0].body
    assert draft.steps[0].images[0].startswith("data:image/png;base64,")
    assert draft.steps[1].title == "Open the recycle bin on the left"
    assert "[Image omitted]" not in draft.overview


def test_troubleshooting_steps_stays_troubleshooting() -> None:
    document = _docx()
    document.add_heading("Troubleshooting steps", level=1)
    document.add_paragraph("Restart the sshd service.")

    draft = parse_docx(_bytes(document), "ssh.docx")

    assert draft.troubleshooting[0].body == "Restart the sshd service."
    assert draft.steps == []


def test_body_only_document_uses_the_filename_as_the_title() -> None:
    document = Document()
    document.add_paragraph("Flush the cache with ipconfig /flushdns.")
    document.add_paragraph("Then confirm with ipconfig /displaydns.")

    draft = parse_docx(_bytes(document), "Flush-DNS-cache.docx")

    assert draft.title == "Flush DNS cache"
    assert draft.summary == "Flush the cache with ipconfig /flushdns."
    assert "displaydns" in draft.overview


def test_a_sentence_that_mentions_a_step_stays_in_the_overview() -> None:
    document = Document()
    document.add_paragraph("First step is to save the excel file from the county.")
    document.add_paragraph("Then click Import.")

    draft = parse_docx(_bytes(document), "wfnj.docx")

    assert draft.steps == []
    assert "save the excel file" in f"{draft.summary} {draft.overview}"


def test_empty_document_is_refused() -> None:
    document = Document()
    document.add_paragraph("   ")

    with pytest.raises(DocxError, match="no text"):
        parse_docx(_bytes(document), "empty.docx")


def test_non_docx_bytes_are_refused() -> None:
    with pytest.raises(DocxError, match="not a Word"):
        parse_docx(b"this is not a zip", "notes.docx")


def test_parse_route_requires_a_bearer_token(client: TestClient) -> None:
    response = client.post(
        "/v1/parse/docx",
        files={"file": ("policy.docx", b"PK\x03\x04not-a-real-docx", "application/octet-stream")},
    )
    assert response.status_code == 401
